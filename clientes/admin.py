# clientes/admin.py
import io
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from django.contrib import admin, messages
from django.http import HttpResponse
from django.urls import path
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from reportlab.lib import colors
from reportlab.platypus import Table, TableStyle
from .models import Cliente


@admin.register(Cliente)
class ClienteAdmin(admin.ModelAdmin):
    list_display = (
        "nombre_completo",
        "tipo_documento",
        "numero_documento",
        "telefono_movil",
        "email",
        "provincia",
        "aplica_descuento",   
        "requiere_envio",     
        "activo",
    )
    list_editable = ("aplica_descuento", "requiere_envio", "activo")  
    search_fields = ("nombre_completo", "numero_documento", "email")
    list_filter = (
        "tipo_documento",
        "provincia",
        "aplica_descuento",  
        "requiere_envio",    
        "activo",
    )
    readonly_fields = ("fecha_creacion", "fecha_actualizacion")
    ordering = ("nombre_completo",)

    actions = ["activar_clientes", "desactivar_clientes"]

    
    change_list_template = "admin/clientes/cliente/change_list.html"

    
    def get_urls(self):
        urls = super().get_urls()
        custom_urls = [
            path(
                "exportar-excel/",
                self.admin_site.admin_view(self.exportar_excel_todo),
                name="clientes_exportar_excel",
            ),
            path(
                "exportar-pdf/",
                self.admin_site.admin_view(self.exportar_pdf_todo),
                name="clientes_exportar_pdf",
            ),
        ]
        return custom_urls + urls

    # -------- Acciones --------
    def activar_clientes(self, request, queryset):
        updated = queryset.update(activo=True)
        self.message_user(
            request, f"✅ {updated} cliente(s) activado(s).", messages.SUCCESS
        )
    activar_clientes.short_description = "Activar clientes seleccionados"

    def desactivar_clientes(self, request, queryset):
        updated = queryset.update(activo=False)
        self.message_user(
            request, f"⚠️ {updated} cliente(s) desactivado(s).", messages.WARNING
        )
    desactivar_clientes.short_description = "Desactivar clientes seleccionados"

    # -------- Exportar Excel --------
    def exportar_excel_todo(self, request):
        queryset = Cliente.objects.all()
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Clientes"

        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill("solid", fgColor="4F81BD")
        center_align = Alignment(horizontal="center", vertical="center")

        headers = [
            "Nombre Completo",
            "Tipo Documento",
            "Número Documento",
            "Teléfono",
            "Email",
            "Provincia",
            "Aplica Descuento",  
            "Requiere Envío",    
            "Activo",
        ]
        ws.append(headers)
        for col_num, _ in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_num)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_align

        for cliente in queryset:
            ws.append(
                [
                    cliente.nombre_completo,
                    cliente.tipo_documento,
                    cliente.numero_documento,
                    cliente.telefono_movil or "",
                    cliente.email or "",
                    cliente.provincia,
                    "Sí" if cliente.aplica_descuento else "No",  # ✅
                    "Sí" if cliente.requiere_envio else "No",    # ✅
                    "Sí" if cliente.activo else "No",
                ]
            )

        for column_cells in ws.columns:
            length = max(len(str(cell.value)) for cell in column_cells) + 2
            ws.column_dimensions[column_cells[0].column_letter].width = length

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        response = HttpResponse(
            output,
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        filename = f"clientes_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        response["Content-Disposition"] = f'attachment; filename="{filename}"'
        return response

    # -------- Exportar PDF --------
    def exportar_pdf_todo(self, request):
        queryset = Cliente.objects.all()
        buffer = io.BytesIO()
        p = canvas.Canvas(buffer, pagesize=A4)
        width, height = A4
        y_start = height - 50

        p.setFont("Helvetica-Bold", 16)
        p.drawCentredString(width / 2, y_start, "REPORTE DE CLIENTES")
        y_start -= 40

        data = [
            [
                "Nombre Completo",
                "Tipo Documento",
                "Número Documento",
                "Teléfono",
                "Email",
                "Provincia",
                "Aplica Desc.",  # ✅
                "Requiere Env.", # ✅
                "Activo",
            ]
        ]

        for cliente in queryset:
            data.append(
                [
                    cliente.nombre_completo,
                    cliente.tipo_documento,
                    cliente.numero_documento,
                    cliente.telefono_movil or "-",
                    cliente.email or "-",
                    cliente.provincia,
                    "Sí" if cliente.aplica_descuento else "No",  # ✅
                    "Sí" if cliente.requiere_envio else "No",    # ✅
                    "Sí" if cliente.activo else "No",
                ]
            )

        table = Table(data, colWidths=[85, 60, 80, 65, 110, 65, 60, 60, 40])
        style = TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#4F81BD")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, 0), 11),
                ("BOTTOMPADDING", (0, 0), (-1, 0), 8),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.black),
            ]
        )
        table.setStyle(style)
        table_width, table_height = table.wrapOn(p, width, height)
        table.drawOn(p, 20, y_start - table_height)

        p.showPage()
        p.save()
        buffer.seek(0)

        response = HttpResponse(buffer, content_type="application/pdf")
        response[
            "Content-Disposition"
        ] = f'attachment; filename="clientes_{datetime.now().strftime("%Y%m%d_%H%M%S")}.pdf"'
        return response
