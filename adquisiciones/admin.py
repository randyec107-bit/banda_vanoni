import io
from datetime import datetime, date, timedelta
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from django.contrib import admin, messages
from django.http import HttpResponse
from django.urls import path
from django.forms.models import BaseInlineFormSet, ModelForm
from django.core.exceptions import ValidationError
from django import forms
from django.db import models
from django.shortcuts import redirect
from .models import Adquisicion, AdquisicionDetalle
from django.utils.html import format_html

# -------- Formularios --------
class AdquisicionForm(ModelForm):
    class Meta:
        model = Adquisicion
        fields = "__all__"
        widgets = {
            'fecha_arribo': forms.DateInput(attrs={'type': 'date'}),
        }

    def clean(self):
        cleaned_data = super().clean()
        return cleaned_data

class AdquisicionDetalleFormSet(BaseInlineFormSet):
    def clean(self):
        super().clean()
        
        # Solo validar si estamos confirmando la adquisición
        if (hasattr(self, 'instance') and 
            self.instance.estado == "CONFIRMADO"):
            
            # Verificar que haya al menos un detalle con datos
            tiene_detalles = False
            for form in self.forms:
                if (not form.cleaned_data.get("DELETE", False) and 
                    form.cleaned_data.get('producto') is not None):
                    tiene_detalles = True
                    break
            
            if not tiene_detalles:
                raise ValidationError(
                    "No se puede confirmar una adquisición sin productos."
                )
            
            # Validar lote y vigencia para cada detalle
            for form in self.forms:
                if (form.cleaned_data.get("DELETE", False) or 
                    form.cleaned_data.get('producto') is None):
                    continue
                    
                lote = form.cleaned_data.get("lote")
                vigencia_lote = form.cleaned_data.get("vigencia_lote")
                
                if not lote or not vigencia_lote:
                    raise ValidationError(
                        "Para confirmar la adquisición, todos los productos deben tener lote y vigencia."
                    )
                
                if vigencia_lote and vigencia_lote <= date.today():
                    form.add_error('vigencia_lote', 'La vigencia del lote debe ser una fecha futura.')

class AdquisicionDetalleInline(admin.TabularInline):
    model = AdquisicionDetalle
    formset = AdquisicionDetalleFormSet
    extra = 1
    fields = (
        "producto",
        "cantidad",
        "lote", 
        "vigencia_lote",
        "peso_unitario_kg",
        "volumen_unitario_m3",
    )
    autocomplete_fields = ["producto"]

# -------- Admin Principal --------
@admin.register(Adquisicion)
class AdquisicionAdmin(admin.ModelAdmin):
    form = AdquisicionForm
    list_display = (
        "numero_orden",
        "proveedor", 
        "fecha_arribo",
        "tipo_carga",
        "estado_display",  
        "inventario_procesado_display",  
        "total_cantidad",
        "total_peso",
        "total_volumen",
    )
    list_filter = ("estado", "tipo_carga", "fecha_arribo", "proveedor", "inventario_procesado")
    search_fields = ("numero_orden", "proveedor__razon_social")
    readonly_fields = ("numero_orden", "inventario_procesado")
    inlines = [AdquisicionDetalleInline]
    
    # formulario de edición
    fieldsets = (
        ('Información Principal', {
            'fields': ('proveedor', 'fecha_arribo', 'tipo_carga', 'estado', 'numero_orden')
        }),
        ('Información del Sistema', {
            'fields': ('inventario_procesado',),
            'classes': ('collapse',),
            'description': 'Estado interno del procesamiento de inventario'
        }),
        ('Información Adicional', {
            'fields': ('observaciones',),
            'classes': ('collapse',)
        }),
    )

    actions = ["accion_confirmar", "accion_cancelar", "exportar_excel_action"]

    def get_form(self, request, obj=None, **kwargs):
        form = super().get_form(request, obj, **kwargs)
        # Establecer fecha mínima en el widget
        form.base_fields['fecha_arribo'].widget.attrs['min'] = date.today().strftime('%Y-%m-%d')
        return form

    
    def estado_display(self, obj):
        color = {
            "BORRADOR": "gray",
            "EN_PROCESO": "blue", 
            "CONFIRMADO": "green",
            "CANCELADO": "red"
        }.get(obj.estado, "black")
        
        return format_html(
            '<span style="color: {}; font-weight: bold;">{}</span>',
            color,
            obj.get_estado_display()
        )
    estado_display.short_description = "Estado"
    
    def inventario_procesado_display(self, obj):
        if obj.inventario_procesado:
            return format_html('✅ <span style="color: green;">Procesado</span>')
        else:
            return format_html('❌ <span style="color: red;">No Procesado</span>')
    inventario_procesado_display.short_description = "Inventario"

    # -------- URLs personalizadas --------
    def get_urls(self):
        urls = super().get_urls()
        custom_urls = [
            path(
                'exportar-excel/',
                self.admin_site.admin_view(self.exportar_excel_todo),
                name='adquisiciones_adquisicion_exportar_excel_todo'
            ),
        ]
        return custom_urls + urls

    
    def save_model(self, request, obj, form, change):
        """Guardar el modelo principal - VERSIÓN MEJORADA"""
        try:
            print(f"🔔 [ADMIN] Iniciando guardado - Estado: {obj.estado}, Cambio: {change}")
            
            
            estado_anterior = None
            if change and obj.pk:
                try:
                    estado_original = Adquisicion.objects.get(pk=obj.pk)
                    estado_anterior = estado_original.estado
                    print(f"🔔 [ADMIN] Estado anterior: {estado_anterior}")
                except Adquisicion.DoesNotExist:
                    pass

            
            super().save_model(request, obj, form, change)
            print(f"🔔 [ADMIN] Objeto guardado - Estado actual: {obj.estado}")

            
            if (estado_anterior != "CONFIRMADO" and 
                obj.estado == "CONFIRMADO" and 
                not obj.inventario_procesado):
                
                print(f"🔔 [ADMIN] Detectado cambio a CONFIRMADO - Ejecutando actualizar_inventario()")
                
                # Ejecutar la actualización del inventario
                try:
                    obj.actualizar_inventario()
                    print(f"🔔 [ADMIN] actualizar_inventario() ejecutado")
                    
                    # Recargar el objeto para ver el estado actualizado
                    obj.refresh_from_db()
                    
                    if obj.inventario_procesado:
                        messages.success(
                            request, 
                            f"✅ Adquisición {obj.numero_orden} CONFIRMADA correctamente. "
                            f"Inventario actualizado con {obj.detalles.count()} productos."
                        )
                    else:
                        messages.warning(
                            request, 
                            f"⚠️ Adquisición {obj.numero_orden} confirmada pero inventario no procesado."
                        )
                        
                except Exception as e:
                    print(f"❌ [ADMIN] Error en actualizar_inventario: {e}")
                    messages.error(
                        request, 
                        f"❌ Error al procesar inventario: {str(e)}"
                    )
                    
            elif obj.estado == "CONFIRMADO" and obj.inventario_procesado:
                messages.info(
                    request, 
                    f"ℹ️ Adquisición {obj.numero_orden} ya estaba confirmada y procesada."
                )
            elif obj.estado == "CANCELADO":
                messages.warning(request, f"⚠️ Adquisición {obj.numero_orden} cancelada.")
            else:
                messages.info(
                    request, 
                    f"💾 Adquisición {obj.numero_orden} guardada en estado {obj.get_estado_display()}."
                )
                
        except Exception as e:
            print(f"❌ [ADMIN] Error general al guardar: {e}")
            messages.error(request, f"❌ Error al guardar: {str(e)}")
            raise

    def save_related(self, request, form, formsets, change):
        """Guardar los inline formsets"""
        try:
            super().save_related(request, form, formsets, change)
            print("🔔 [ADMIN] Detalles guardados correctamente")
        except Exception as e:
            print(f"❌ [ADMIN] Error en detalles: {e}")
            messages.error(request, f"❌ Error en los detalles: {str(e)}")
            raise

    
    def accion_confirmar(self, request, queryset):
        """Action para confirmar múltiples adquisiciones - VERSIÓN CORREGIDA"""
        success_count = 0
        error_count = 0
        
        for adquisicion in queryset:
            try:
                print(f"🔔 [ACTION] Procesando {adquisicion.numero_orden}")
                
                # Verificar que no esté ya confirmada
                if adquisicion.estado == "CONFIRMADO":
                    self.message_user(
                        request, 
                        f"ℹ️ {adquisicion.numero_orden}: Ya está confirmada.", 
                        messages.INFO
                    )
                    continue
                
                # Verificar que tenga detalles
                if not adquisicion.detalles.exists():
                    self.message_user(
                        request, 
                        f"❌ {adquisicion.numero_orden}: No se puede confirmar sin detalles.", 
                        messages.ERROR
                    )
                    error_count += 1
                    continue
                
                # Verificar detalles completos
                detalles_incompletos = adquisicion.detalles.filter(
                    models.Q(lote__isnull=True) | 
                    models.Q(lote='') |
                    models.Q(vigencia_lote__isnull=True)
                )
                
                if detalles_incompletos.exists():
                    self.message_user(
                        request, 
                        f"❌ {adquisicion.numero_orden}: Tiene productos sin lote o vigencia.", 
                        messages.ERROR
                    )
                    error_count += 1
                    continue
                
                # Verificar vigencia
                detalles_vigencia_invalida = adquisicion.detalles.filter(
                    vigencia_lote__lte=date.today()
                )
                
                if detalles_vigencia_invalida.exists():
                    self.message_user(
                        request, 
                        f"❌ {adquisicion.numero_orden}: Tiene productos con vigencia vencida.", 
                        messages.ERROR
                    )
                    error_count += 1
                    continue
                
                
                adquisicion.estado = "CONFIRMADO"
                adquisicion.save()  
                
                # Recargar para ver estado actualizado
                adquisicion.refresh_from_db()
                
                if adquisicion.inventario_procesado:
                    success_count += 1
                    self.message_user(
                        request, 
                        f"✅ {adquisicion.numero_orden} confirmada e inventario actualizado.", 
                        messages.SUCCESS
                    )
                else:
                    error_count += 1
                    self.message_user(
                        request, 
                        f"⚠️ {adquisicion.numero_orden} confirmada pero inventario no procesado.", 
                        messages.WARNING
                    )
                
            except Exception as e:
                error_count += 1
                self.message_user(
                    request, 
                    f"❌ Error en {adquisicion.numero_orden}: {str(e)}", 
                    messages.ERROR
                )
        
        # Resumen final
        if success_count > 0:
            self.message_user(
                request, 
                f"🎉 {success_count} adquisiciones confirmadas exitosamente.", 
                messages.SUCCESS
            )
        if error_count > 0:
            self.message_user(
                request, 
                f"❌ {error_count} adquisiciones tuvieron problemas.", 
                messages.ERROR
            )
            
    accion_confirmar.short_description = "✅ Confirmar adquisiciones seleccionadas"

    def accion_cancelar(self, request, queryset):
        """Action para cancelar múltiples adquisiciones"""
        success_count = 0
        for adquisicion in queryset:
            try:
                if adquisicion.estado != "CANCELADO":
                    adquisicion.estado = "CANCELADO"
                    adquisicion.save()
                    success_count += 1
                    self.message_user(
                        request, 
                        f"⚠️ {adquisicion.numero_orden} cancelada.", 
                        messages.WARNING
                    )
                else:
                    self.message_user(
                        request, 
                        f"ℹ️ {adquisicion.numero_orden} ya estaba cancelada.", 
                        messages.INFO
                    )
            except Exception as e:
                self.message_user(
                    request, 
                    f"❌ Error cancelando {adquisicion.numero_orden}: {str(e)}", 
                    messages.ERROR
                )
        
        if success_count > 0:
            self.message_user(
                request, 
                f"🎯 {success_count} adquisiciones canceladas.", 
                messages.SUCCESS
            )
            
    accion_cancelar.short_description = "❌ Cancelar adquisiciones seleccionadas"

    # -------- Action para exportar Excel --------
    def exportar_excel_action(self, request, queryset):
        """Action para exportar adquisiciones seleccionadas a Excel"""
        try:
            wb = openpyxl.Workbook()
            ws_resumen = wb.active
            ws_resumen.title = "Resumen Adquisiciones"

            
            header_font = Font(bold=True, color="FFFFFF")
            header_fill = PatternFill("solid", fgColor="4F81BD")
            center_align = Alignment(horizontal="center", vertical="center")

            
            headers = [
                "N° Orden", "Proveedor", "Fecha Arribo", "Tipo Carga",
                "Estado", "Inventario Procesado", "Total Cantidad", "Total Peso (kg)", "Total Volumen (m3)"
            ]
            ws_resumen.append(headers)
            
            
            for col_num in range(1, len(headers) + 1):
                cell = ws_resumen.cell(row=1, column=col_num)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = center_align

            
            for adquisicion in queryset:
                ws_resumen.append([
                    adquisicion.numero_orden,
                    str(adquisicion.proveedor),
                    adquisicion.fecha_arribo.strftime("%d/%m/%Y"),
                    adquisicion.get_tipo_carga_display(),
                    adquisicion.get_estado_display(),
                    "SÍ" if adquisicion.inventario_procesado else "NO",
                    adquisicion.total_cantidad,
                    adquisicion.total_peso or 0,
                    adquisicion.total_volumen or 0,
                ])

            
            for column_cells in ws_resumen.columns:
                length = max(len(str(cell.value or "")) for cell in column_cells) + 2
                ws_resumen.column_dimensions[column_cells[0].column_letter].width = min(length, 50)

            # Guardar en buffer
            output = io.BytesIO()
            wb.save(output)
            output.seek(0)
            
            # Crear respuesta
            response = HttpResponse(
                output, 
                content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )
            filename = f"adquisiciones_seleccionadas_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
            response["Content-Disposition"] = f'attachment; filename="{filename}"'
            
            messages.success(request, f"✅ Excel exportado correctamente con {queryset.count()} adquisiciones.")
            return response
            
        except Exception as e:
            messages.error(request, f"❌ Error al exportar Excel: {str(e)}")
            return redirect('admin:adquisiciones_adquisicion_changelist')
            
    exportar_excel_action.short_description = "📊 Exportar adquisiciones seleccionadas a Excel"

    # -------- Exportar Excel para todas las adquisiciones --------
    def exportar_excel_todo(self, request):
        """Vista para exportar TODAS las adquisiciones a Excel"""
        try:
            queryset = Adquisicion.objects.all().select_related('proveedor').prefetch_related('detalles__producto')
            wb = openpyxl.Workbook()
            ws_resumen = wb.active
            ws_resumen.title = "Resumen Adquisiciones"

            # Estilos
            header_font = Font(bold=True, color="FFFFFF")
            header_fill = PatternFill("solid", fgColor="4F81BD")
            center_align = Alignment(horizontal="center", vertical="center")

            
            headers = [
                "N° Orden", "Proveedor", "Fecha Arribo", "Tipo Carga",
                "Estado", "Inventario Procesado", "Total Cantidad", "Total Peso (kg)", "Total Volumen (m3)"
            ]
            ws_resumen.append(headers)
            
            # Aplicar estilos a los encabezados
            for col_num in range(1, len(headers) + 1):
                cell = ws_resumen.cell(row=1, column=col_num)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = center_align

            
            for adquisicion in queryset:
                ws_resumen.append([
                    adquisicion.numero_orden,
                    str(adquisicion.proveedor),
                    adquisicion.fecha_arribo.strftime("%d/%m/%Y"),
                    adquisicion.get_tipo_carga_display(),
                    adquisicion.get_estado_display(),
                    "SÍ" if adquisicion.inventario_procesado else "NO",
                    adquisicion.total_cantidad,
                    adquisicion.total_peso or 0,
                    adquisicion.total_volumen or 0,
                ])

            # Ajustar ancho de columnas en el resumen
            for column_cells in ws_resumen.columns:
                length = max(len(str(cell.value or "")) for cell in column_cells) + 2
                ws_resumen.column_dimensions[column_cells[0].column_letter].width = min(length, 50)

            # Hojas individuales por adquisición
            for adquisicion in queryset:
                # Limitar el nombre de la hoja a 31 caracteres (límite de Excel)
                sheet_name = adquisicion.numero_orden[:31]
                # Evitar nombres duplicados
                if sheet_name in wb.sheetnames:
                    sheet_name = f"{sheet_name}_{adquisicion.id}"
                
                ws = wb.create_sheet(title=sheet_name[:31])
                
                
                ws_headers = [
                    "Producto", "Cantidad", "Lote", "Vigencia Lote",
                    "Peso Unitario (kg)", "Volumen Unitario (m3)", "Peso Total", "Volumen Total"
                ]
                ws.append(ws_headers)
                
                
                for col_num in range(1, len(ws_headers) + 1):
                    cell = ws.cell(row=1, column=col_num)
                    cell.font = header_font
                    cell.fill = header_fill
                    cell.alignment = center_align

                # Datos de detalles
                for detalle in adquisicion.detalles.all():
                    ws.append([
                        detalle.producto.nombre if detalle.producto else "",
                        detalle.cantidad,
                        detalle.lote or "",
                        detalle.vigencia_lote.strftime("%d/%m/%Y") if detalle.vigencia_lote else "",
                        float(detalle.peso_unitario_kg) if detalle.peso_unitario_kg else 0,
                        float(detalle.volumen_unitario_m3) if detalle.volumen_unitario_m3 else 0,
                        float(detalle.peso_total) if detalle.peso_total else 0,
                        float(detalle.volumen_total) if detalle.volumen_total else 0
                    ])

                
                for column_cells in ws.columns:
                    length = max(len(str(cell.value or "")) for cell in column_cells) + 2
                    ws.column_dimensions[column_cells[0].column_letter].width = min(length, 30)

            # Guardar en buffer
            output = io.BytesIO()
            wb.save(output)
            output.seek(0)
            
            # Crear respuesta
            response = HttpResponse(
                output, 
                content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )
            filename = f"adquisiciones_completo_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
            response["Content-Disposition"] = f'attachment; filename="{filename}"'
            
            messages.success(request, f"✅ Excel exportado correctamente con {queryset.count()} adquisiciones.")
            return response
            
        except Exception as e:
            messages.error(request, f"❌ Error al exportar Excel: {str(e)}")
            return redirect('admin:adquisiciones_adquisicion_changelist')

    # -------- Cambios en la respuesta después de guardar --------
    def response_change(self, request, obj):
        """Manejar la respuesta después de guardar un objeto existente"""
        response = super().response_change(request, obj)
        
        # Mensaje especial para confirmaciones
        if obj.estado == "CONFIRMADO" and '_save' in request.POST:
            if obj.inventario_procesado:
                messages.success(request, f"✅ Adquisición {obj.numero_orden} confirmada e inventario actualizado!")
            else:
                messages.warning(request, f"⚠️ Adquisición {obj.numero_orden} confirmada pero inventario no procesado.")
            
        return response

    def response_add(self, request, obj, post_url_continue=None):
        """Manejar la respuesta después de agregar un nuevo objeto"""
        response = super().response_add(request, obj, post_url_continue)
        
        if obj.estado == "CONFIRMADO":
            if obj.inventario_procesado:
                messages.success(request, f"✅ Adquisición {obj.numero_orden} confirmada e inventario actualizado!")
            else:
                messages.warning(request, f"⚠️ Adquisición {obj.numero_orden} confirmada pero inventario no procesado.")
            
        return response

    # -------- Campos de solo lectura para estados específicos --------
    def get_readonly_fields(self, request, obj=None):
        readonly_fields = list(super().get_readonly_fields(request, obj))
        if obj and obj.estado == "CONFIRMADO":
            readonly_fields.extend(['proveedor', 'fecha_arribo', 'tipo_carga'])
        return readonly_fields

    # -------- Permisos para acciones --------
    def get_actions(self, request):
        actions = super().get_actions(request)
        if not request.user.has_perm('adquisiciones.change_adquisicion'):
            # Remover acciones si no tiene permiso de cambio
            if 'accion_confirmar' in actions:
                del actions['accion_confirmar']
            if 'accion_cancelar' in actions:
                del actions['accion_cancelar']
            if 'exportar_excel_action' in actions:
                del actions['exportar_excel_action']
        return actions