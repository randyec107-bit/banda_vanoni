# adquisiciones/views.py
import io
import openpyxl
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.contrib import messages
from django.contrib.auth.decorators import login_required, permission_required
from django.views.decorators.http import require_POST
from .models import Adquisicion, AdquisicionDetalle

def adquisiciones_home(request):
    """
    Página principal del módulo de adquisiciones
    """
    # Estadísticas básicas
    total_adquisiciones = Adquisicion.objects.count()
    adquisiciones_recientes = Adquisicion.objects.all().order_by('-id')[:5]
    
    # Contar por estado
    estados = {
        'BORRADOR': Adquisicion.objects.filter(estado="BORRADOR").count(),
        'EN_PROCESO': Adquisicion.objects.filter(estado="EN_PROCESO").count(),
        'CONFIRMADO': Adquisicion.objects.filter(estado="CONFIRMADO").count(),
        'CANCELADO': Adquisicion.objects.filter(estado="CANCELADO").count(),
    }
    
    context = {
        'total_adquisiciones': total_adquisiciones,
        'adquisiciones_recientes': adquisiciones_recientes,
        'estados': estados,
    }
    return render(request, 'adquisiciones/home.html', context)

def exportar_adquisiciones_excel(request):
    """
    Exportar adquisiciones seleccionadas a Excel
    Recibe IDs por parámetro: /exportar_excel/?ids=1&ids=2&ids=3
    """
    try:
        
        ids = request.GET.getlist("ids")
        
        if not ids:
            # Si no hay IDs, exportar todas las adquisiciones
            adquisiciones = Adquisicion.objects.all().select_related('proveedor').prefetch_related('detalles__producto')
            filename_suffix = "completo"
        else:
            # Exportar solo las adquisiciones seleccionadas
            adquisiciones = Adquisicion.objects.filter(id__in=ids).select_related('proveedor').prefetch_related('detalles__producto')
            filename_suffix = f"seleccionadas_{len(ids)}"

        # Crear workbook
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Resumen Adquisiciones"

        
        headers = [
            "N° Orden", 
            "Proveedor", 
            "Fecha Arribo", 
            "Tipo Carga",
            "Estado", 
            "Total Cantidad", 
            "Total Peso (kg)", 
            "Total Volumen (m3)", 
            "Inventario Procesado",
            "Productos Únicos"
        ]
        ws.append(headers)

        
        for col in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col)
            cell.font = openpyxl.styles.Font(bold=True, color="FFFFFF")
            cell.fill = openpyxl.styles.PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")

        
        for adquisicion in adquisiciones:
            ws.append([
                adquisicion.numero_orden,
                str(adquisicion.proveedor),
                adquisicion.fecha_arribo.strftime("%d/%m/%Y"),
                adquisicion.get_tipo_carga_display(),
                adquisicion.get_estado_display(),
                adquisicion.total_cantidad,
                float(adquisicion.total_peso) if adquisicion.total_peso else 0,
                float(adquisicion.total_volumen) if adquisicion.total_volumen else 0,
                "✅ Sí" if adquisicion.inventario_procesado else "❌ No",
                adquisicion.productos_unicos
            ])

        
        column_widths = {
            'A': 15,  # N° Orden
            'B': 30,  # Proveedor
            'C': 12,  # Fecha Arribo
            'D': 15,  # Tipo Carga
            'E': 12,  # Estado
            'F': 15,  # Total Cantidad
            'G': 15,  # Total Peso
            'H': 15,  # Total Volumen
            'I': 18,  # Inventario Procesado
            'J': 15,  # Productos Únicos
        }

        for col, width in column_widths.items():
            ws.column_dimensions[col].width = width

        # Generar respuesta
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        response = HttpResponse(
            output,
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        filename = f"adquisiciones_{filename_suffix}_{adquisiciones.count()}.xlsx"
        response["Content-Disposition"] = f'attachment; filename="{filename}"'
        
        return response

    except Exception as e:
        # En caso de error, redirigir a la página principal con mensaje
        messages.error(request, f"❌ Error al exportar Excel: {str(e)}")
        return redirect('adquisiciones_home')

@login_required
@permission_required('adquisiciones.change_adquisicion', raise_exception=True)
@require_POST
def confirmar_adquisicion(request, adquisicion_id):
    """
    Vista para confirmar una adquisición y procesar automáticamente el inventario
    """
    adquisicion = get_object_or_404(Adquisicion, id=adquisicion_id)
    
    # Verificar si ya está confirmada
    if adquisicion.estado == "CONFIRMADO":
        messages.warning(request, f'La adquisición {adquisicion.numero_orden} ya está confirmada.')
        return redirect('admin:adquisiciones_adquisicion_change', object_id=adquisicion_id)
    
    
    puede_confirmar, mensaje = adquisicion.puede_confirmar()
    
    if not puede_confirmar:
        messages.error(request, f'No se puede confirmar la adquisición: {mensaje}')
        return redirect('admin:adquisiciones_adquisicion_change', object_id=adquisicion_id)
    
    try:
        
        if adquisicion.confirmar():
            messages.success(
                request, 
                f'✅ Adquisición {adquisicion.numero_orden} confirmada exitosamente. '
                f'Se han creado/actualizado {adquisicion.detalles.count()} lotes en el inventario.'
            )
        else:
            messages.warning(request, f'La adquisición {adquisicion.numero_orden} no pudo ser confirmada.')
        
    except Exception as e:
        messages.error(
            request, 
            f'❌ Error al confirmar la adquisición: {str(e)}'
        )
    
    return redirect('admin:adquisiciones_adquisicion_change', object_id=adquisicion_id)

@login_required
def lista_adquisiciones(request):
    """
    Vista para listar todas las adquisiciones (opcional)
    """
    estado_filter = request.GET.get('estado', '')
    
    adquisiciones = Adquisicion.objects.all().select_related('proveedor').order_by('-fecha_arribo')
    
    if estado_filter:
        adquisiciones = adquisiciones.filter(estado=estado_filter)
    
    # Estadísticas
    estadisticas = {
        'total': Adquisicion.objects.count(),
        'borrador': Adquisicion.objects.filter(estado="BORRADOR").count(),
        'en_proceso': Adquisicion.objects.filter(estado="EN_PROCESO").count(),
        'confirmado': Adquisicion.objects.filter(estado="CONFIRMADO").count(),
        'cancelado': Adquisicion.objects.filter(estado="CANCELADO").count(),
    }
    
    context = {
        'adquisiciones': adquisiciones,
        'estadisticas': estadisticas,
        'estado_filter': estado_filter,
        'estado_choices': Adquisicion.ESTADO_CHOICES,
    }
    
    return render(request, 'adquisiciones/lista_adquisiciones.html', context)

@login_required
def detalle_adquisicion(request, adquisicion_id):
    """
    Vista para ver el detalle de una adquisición (opcional)
    """
    adquisicion = get_object_or_404(
        Adquisicion.objects.prefetch_related('detalles__producto'), 
        id=adquisicion_id
    )
    
    context = {
        'adquisicion': adquisicion,
        'detalles': adquisicion.detalles.all(),
        'puede_confirmar': adquisicion.estado in ['BORRADOR', 'EN_PROCESO'],
    }
    
    return render(request, 'adquisiciones/detalle_adquisicion.html', context)